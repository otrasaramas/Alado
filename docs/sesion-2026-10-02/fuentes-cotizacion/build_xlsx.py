import sys
from datetime import date
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
sys.path.insert(0, ".")
from data import RATES, ITEMS, OPTIONAL

OUT = sys.argv[1]
wb = Workbook(); ws = wb.active; ws.title = "Cotización"
INK = "2E3A46"; BLUE = "5C7891"; PAPER = "F3F1EA"
YEL = PatternFill("solid", fgColor="FFF2B3"); HEAD = PatternFill("solid", fgColor=BLUE); SEC = PatternFill("solid", fgColor="E6EAEE"); TOT = PatternFill("solid", fgColor=PAPER)
thin = Side(style="thin", color="C9D2DB"); B = Border(bottom=thin)
money = '"$"#,##0'; datefmt = 'DD/MM/YYYY'
widths = [5, 13, 30, 46, 18, 12, 16, 11, 7, 13, 14, 38]
for i, w in enumerate(widths, 1): ws.column_dimensions[get_column_letter(i)].width = w

ws["A1"] = "Cotización de contenido · Alado"; ws["A1"].font = Font(size=16, bold=True, color=INK)
ws["A2"] = "Alado Bride y Alado & Co. · octubre 2026 – enero 2027 · Preparada por Sara · 14 de octubre de 2026"; ws["A2"].font = Font(size=10, color="5A6672")

ws["A4"] = "TARIFAS · edita las celdas amarillas"; ws["A4"].font = Font(bold=True, color=BLUE)
r = 5; rate_first = r
for fmt, val, note in RATES:
    ws.cell(r, 2, fmt).font = Font(bold=True, color=INK)
    c = ws.cell(r, 3, val); c.fill = YEL; c.number_format = money; c.font = Font(bold=True, color="1F4E99")
    ws.cell(r, 4, note).font = Font(italic=True, color="5A6672")
    r += 1
rate_last = r - 1
RT = f"$B${rate_first}:$C${rate_last}"

r += 1
hdr = ["#", "Línea", "Proyecto", "Pieza", "Formato", "Graba", "Publica", "Mes", "Cant.", "Tarifa", "Subtotal", "Nota"]
for j, h in enumerate(hdr, 1):
    c = ws.cell(r, j, h); c.fill = HEAD; c.font = Font(bold=True, color="FFFFFF"); c.alignment = Alignment(vertical="center")
hdr_row = r; r += 1
first_item = r
MONTHS = '"Ene","Feb","Mar","Abr","May","Jun","Jul","Ago","Sep","Oct","Nov","Dic"'
n = 0
def item_row(row, it, idx, qty_input=False):
    linea, proy, pieza, fmt, graba, pub, qty, note = it
    ws.cell(row, 1, idx); ws.cell(row, 2, linea); ws.cell(row, 3, proy); ws.cell(row, 4, pieza); ws.cell(row, 5, fmt)
    if graba: c = ws.cell(row, 6, graba); c.number_format = datefmt
    if pub == "asap": ws.cell(row, 7, "Lo antes posible")
    elif pub: c = ws.cell(row, 7, pub); c.number_format = datefmt
    ws.cell(row, 8, f'=IF(ISNUMBER(G{row}),CHOOSE(MONTH(G{row}),{MONTHS}),IF(G{row}="Lo antes posible","Oct","Por definir"))')
    c = ws.cell(row, 9, qty)
    if qty_input or qty != 1 or "confirmar" in note.lower(): c.fill = YEL; c.font = Font(bold=True, color="1F4E99")
    c = ws.cell(row, 10, f'=IF(VLOOKUP(E{row},{RT},2,FALSE)="","",VLOOKUP(E{row},{RT},2,FALSE))'); c.number_format = money
    c = ws.cell(row, 11, f'=IF(J{row}="","",I{row}*J{row})'); c.number_format = money
    ws.cell(row, 12, note).font = Font(italic=True, color="5A6672")
    for j in range(1, 13): ws.cell(row, j).border = B

sections = {}
def section(title, items, qty_input=False):
    global r, n
    c = ws.cell(r, 1, title); c.font = Font(bold=True, color=INK)
    for j in range(1, 13): ws.cell(r, j).fill = SEC
    r += 1; start = r
    for it in items:
        n += 1; item_row(r, it, n, qty_input); r += 1
    end = r - 1
    ws.cell(r, 10, f"Subtotal {title.split(' · ')[0].lower()}").font = Font(bold=True)
    c = ws.cell(r, 11, f"=SUM(K{start}:K{end})"); c.number_format = money; c.font = Font(bold=True)
    sections[title] = (start, end, r); r += 2

section("ALADO BRIDE", [i for i in ITEMS if i[0] == "Bride"])
section("ALADO & CO.", [i for i in ITEMS if i[0] == "Alado & Co."])
section("OPCIONALES DEL TALLER DE CERÁMICA · cantidad 0, cambia a 1 si se cobran aparte", OPTIONAL, qty_input=True)
last_item = r - 3

# Resumen
ws.cell(r, 1, "RESUMEN").font = Font(bold=True, color=BLUE); r += 1
sb = sections["ALADO BRIDE"][2]; sc = sections["ALADO & CO."][2]; so = sections["OPCIONALES DEL TALLER DE CERÁMICA · cantidad 0, cambia a 1 si se cobran aparte"][2]
summ = [("Alado Bride", f"=K{sb}"), ("Alado & Co.", f"=K{sc}"), ("Opcionales", f"=K{so}")]
sum_first = r
for label, f in summ:
    ws.cell(r, 4, label); c = ws.cell(r, 11, f); c.number_format = money; r += 1
ws.cell(r, 4, "TOTAL").font = Font(bold=True, size=13)
c = ws.cell(r, 11, f"=SUM(K{sum_first}:K{r-1})"); c.number_format = money; c.font = Font(bold=True, size=13)
for j in range(1, 13): ws.cell(r, j).fill = TOT
total_row = r; r += 1
ws.cell(r, 4, "Gestión de talleres y campaña de fotos todavía no tienen tarifa: el total sube cuando las definas.").font = Font(italic=True, color="9C5700")
r += 2

ws.cell(r, 1, "POR MES (según fecha de publicación)").font = Font(bold=True, color=BLUE); r += 1
rng = lambda col: f"{col}{first_item}:{col}{last_item}"
month_first = r
for m in ["Oct", "Nov", "Dic", "Ene", "Por definir"]:
    ws.cell(r, 4, m); c = ws.cell(r, 11, f'=SUMIF({rng("H")},"{m}",{rng("K")})'); c.number_format = money; r += 1
ws.cell(r, 4, "Total por mes (debe ser igual al TOTAL)").font = Font(italic=True, color="5A6672")
c = ws.cell(r, 11, f"=SUM(K{month_first}:K{r-1})"); c.number_format = money; r += 2

ws.cell(r, 1, "CONDICIONES").font = Font(bold=True, color=BLUE); r += 1
for t in ["50% al aprobar y 50% contra entrega, por mes.",
          "Hasta dos rondas de ajustes por pieza.",
          "Piezas extra a la tarifa vigente.",
          "Gastos reembolsables aparte, previa aprobación (transporte, equipo, música con licencia).",
          "Los derechos del contenido son de Alado; Sara puede mostrarlo en su portafolio."]:
    ws.cell(r, 2, "· " + t); r += 1
r += 1
ws.cell(r, 1, "DATOS QUE FALTAN").font = Font(bold=True, color="9C5700"); r += 1
for t in ["Tarifa de gestión y organización por taller.",
          "Tarifa de la campaña de fotos de textil hogar (y cuántas fotos editadas se entregan).",
          "Tarifa de pieza gráfica suelta (hoy igual al set de historias).",
          "Número de partes de la serie con Lilo.",
          "Fechas: servicio con Lola, Alejo como artista, Alejo 7 claves, video promocional del kimono, taller de los Pérez.",
          "¿Salieron el Destacado «Bride» (9 oct) y el figurín en vivo (13 oct)?",
          "¿Se cotizan el video de los Pérez y el contenido del lanzamiento con Entrelazos (27 nov)?",
          "¿Los opcionales del taller de cerámica van dentro de gestión o aparte?"]:
    ws.cell(r, 2, "· " + t); r += 1

ws.freeze_panes = ws.cell(hdr_row + 1, 1)
ws.sheet_view.showGridLines = False
wb.save(OUT)
print("total row", total_row, "items", first_item, last_item)
