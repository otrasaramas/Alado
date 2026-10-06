import json, re
from pathlib import Path
from pptx import Presentation
from pptx.util import Emu, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
OUT = Path("out2"); E = lambda v: Emu(int(round(v*9525)))
def rgb(css):
    r,g,b = [int(x) for x in re.findall(r"\d+", css)[:3]]; return RGBColor(r,g,b)
FAM = lambda f, w: "Archivo Black" if f == "Archivo Black" else {500:"Archivo Medium",600:"Archivo SemiBold",700:"Archivo"}.get(w,"Archivo")
meas = json.loads((OUT/"medidas.json").read_text())
prs = Presentation(); prs.slide_width, prs.slide_height = E(1080), E(1920)
for name, items in meas.items():
    s = prs.slides.add_slide(prs.slide_layouts[6])
    s.background.fill.solid(); s.background.fill.fore_color.rgb = RGBColor(0x5C,0x78,0x91)
    for lay, label in (("fondo","Foto de fondo (reemplazar)"),("ticket","Ticket"),("sello","Sello")):
        pic = s.shapes.add_picture(str(OUT/f"capa-{lay}-{name}.png"), 0, 0, E(1080), E(1920)); pic.name = label
    for d in items:
        x, w = d["x"], d["w"]
        if d["align"] == "left": w += 40
        tb = s.shapes.add_textbox(E(x), E(d["y"]), E(w), E(d["h"]))
        tf = tb.text_frame; tf.word_wrap = False; tf.auto_size = None; tf.vertical_anchor = MSO_ANCHOR.TOP
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
        for i, line in enumerate(d["lines"]):
            para = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            para.alignment = {"center":PP_ALIGN.CENTER,"right":PP_ALIGN.RIGHT}.get(d["align"], PP_ALIGN.LEFT)
            para.line_spacing = Pt(d["lh"]*0.75)
            r = para.add_run(); r.text = line; f = r.font
            f.name = FAM(d["family"], d["weight"]); f.size = Pt(d["size"]*0.75); f.color.rgb = rgb(d["color"])
            f.bold = d["weight"] == 700 and d["family"] != "Archivo Black"
            if d["spacing"]: r._r.get_or_add_rPr().set("spc", str(int(round(d["spacing"]*75))))
prs.save(OUT/"historias-taller-ceramica.pptx"); print("ok")
