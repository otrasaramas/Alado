import json, re
from pathlib import Path
from pptx import Presentation
from pptx.util import Emu, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from gen import THEMES, VX, VY, VW, VH

OUT = Path("out"); PX = 9525
def E(v): return Emu(int(round(v*PX)))
def rgb(css):
    if css.startswith("#"): return RGBColor.from_string(css[1:].upper())
    r,g,b = [int(x) for x in re.findall(r"\d+", css)[:3]]; return RGBColor(r,g,b)
NAMES = {"t-labels":"Datos (etiquetas)","t-values":"Datos (valores)","t-title1":"Título 1","t-tag":"Tagline",
         "t-title2":"Título 2","t-left":"Abajo izquierda","t-right":"Abajo derecha"}

for th, t in THEMES.items():
    prs = Presentation(); prs.slide_width, prs.slide_height = E(1080), E(1920)
    s = prs.slides.add_slide(prs.slide_layouts[6])
    s.background.fill.solid(); s.background.fill.fore_color.rgb = rgb(t["bg"])
    ph = s.shapes.add_picture(str(OUT/f"capa-foto-ejemplo-{th}.png"), E(VX), E(VY), E(VW), E(VH)); ph.name = "FOTO (reemplazar)"
    win = s.shapes.add_picture(str(OUT/f"capa-fondo-ventana-{th}.png"), 0, 0, E(1080), E(1920)); win.name = "Fondo con ventana"
    for d in json.loads((OUT/f"medidas-{th}.json").read_text()):
        right = d["align"] in ("right","end"); pad = 30
        x = d["x"] - (pad if right else 0); w = d["w"] + pad
        tb = s.shapes.add_textbox(E(x), E(d["y"]), E(w), E(d["h"])); tb.name = NAMES[d["id"]]
        tf = tb.text_frame; tf.word_wrap = False; tf.auto_size = None; tf.vertical_anchor = MSO_ANCHOR.TOP
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
        for i, line in enumerate(d["lines"]):
            para = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            para.alignment = PP_ALIGN.RIGHT if right else PP_ALIGN.LEFT
            para.line_spacing = Pt(d["lh"]*0.75)
            r = para.add_run(); r.text = line
            f = r.font; f.name = d["family"]; f.size = Pt(d["size"]*0.75); f.italic = d["italic"]
            f.color.rgb = rgb(d["color"])
            if d["spacing"]: r._r.get_or_add_rPr().set("spc", str(int(round(d["spacing"]*0.75*100))))
    prs.save(OUT/f"historia-ventana-{th}.pptx"); print("ok", th)
